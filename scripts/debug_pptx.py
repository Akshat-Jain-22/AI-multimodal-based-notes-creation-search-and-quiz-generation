from pptx import Presentation
import sys

def debug_slide_shapes(pptx_path, slide_numbers):
    prs = Presentation(pptx_path)
    for i, slide in enumerate(prs.slides, start=1):
        if i in slide_numbers:
            print(f"\n=== Slide {i} ===")
            print(f"Number of shapes: {len(slide.shapes)}")
            for j, shape in enumerate(slide.shapes):
                print(f"  Shape {j}: type={shape.shape_type}, name='{shape.name}'")

if __name__ == "__main__":
    pptx_file = sys.argv[1]
    debug_slide_shapes(pptx_file, [6, 9, 10])