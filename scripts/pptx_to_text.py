from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image
import pytesseract
import io
import subprocess
import os
import sys
import json
from pdf2image import convert_from_path


def process_shapes(shapes, slide_text, image_entries, slide_num, images_dir):
    """Recursively walk through shapes (including grouped ones), extracting
    text and saving any embedded pictures as actual PNG files."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            process_shapes(shape.shapes, slide_text, image_entries, slide_num, images_dir)
            continue

        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if text:
                    slide_text.append(text)

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                image_bytes = shape.image.blob
                pil_img = Image.open(io.BytesIO(image_bytes))

                img_filename = f"slide{slide_num}_img{len(image_entries) + 1}.png"
                img_path = os.path.join(images_dir, img_filename)
                pil_img.save(img_path)

                ocr_text = pytesseract.image_to_string(pil_img).strip()

                image_entries.append({
                    "source": "embedded_picture",
                    "image_file": img_filename,   
                    "ocr_text": ocr_text
                })
            except Exception as e:
                print(f"  Warning: could not process an image on slide {slide_num}: {e}")


def convert_pptx_to_pdf(pptx_path, output_dir="temp_pdf"):
    os.makedirs(output_dir, exist_ok=True)
    subprocess.run([
        "soffice", "--headless", "--convert-to", "pdf",
        "--outdir", output_dir, pptx_path
    ], check=True)
    pdf_name = os.path.splitext(os.path.basename(pptx_path))[0] + ".pdf"
    return os.path.join(output_dir, pdf_name)


def extract_pptx_content(pptx_path):
    prs = Presentation(pptx_path)

    pptx_basename = os.path.splitext(os.path.basename(pptx_path))[0]
    images_dir = os.path.join("extracted_images", pptx_basename)
    os.makedirs(images_dir, exist_ok=True)

    structured_slides = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_text = []
        image_entries = []
        process_shapes(slide.shapes, slide_text, image_entries, i, images_dir)

        notes_text = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()

        structured_slides.append({
            "slide_number": i,
            "content": slide_text,
            "images": image_entries,     
            "speaker_notes": notes_text
        })

    print("Rendering full slides as fallback for visual content (diagrams, freeform shapes)...")
    pdf_path = convert_pptx_to_pdf(pptx_path)
    pages = convert_from_path(pdf_path, dpi=200)

    for i, page_image in enumerate(pages, start=1):
        slide_data = structured_slides[i - 1]
        has_content = bool(slide_data["content"]) or bool(slide_data["images"])

        if not has_content:
            full_img_filename = f"slide{i}_full.png"
            full_img_path = os.path.join(images_dir, full_img_filename)
            page_image.save(full_img_path)

            full_slide_text = pytesseract.image_to_string(page_image).strip()

            slide_data["images"].append({
                "source": "full_slide_render",
                "image_file": full_img_filename,
                "ocr_text": full_slide_text
            })
            print(f"Slide {i}: saved full-slide image ({full_img_filename}) + OCR ({len(full_slide_text)} chars)")
        else:
            print(f"Slide {i}: structured extraction sufficient ({len(slide_data['content'])} text blocks)")

    return {
        "source_pptx": os.path.basename(pptx_path),
        "images_directory": images_dir,
        "slides": structured_slides
    }


if __name__ == "__main__":
    pptx_file = sys.argv[1] if len(sys.argv) > 1 else "sample_slides.pptx"
    result = extract_pptx_content(pptx_file)

    with open("pptx_content.json", "w", encoding="utf-8") as f:
        json.dump(result, f, indent=2, ensure_ascii=False)

    print(f"\nExtracted {len(result['slides'])} slides to pptx_content.json")
    print(f"Images saved to: {result['images_directory']}/")